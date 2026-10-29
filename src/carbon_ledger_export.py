# -*- coding: utf-8 -*-
"""碳账交付（carbon_ledger/）→ 前端只读结果表：适配器。

定位
----
这个模块**只搬运 + 预计算**，不重算碳账、不改交付方任何文件：

* 它读 ``carbon_ledger/`` 下的 A2/A4/A5 交付工作簿（Excel），把里面**已经算好的**
  值与判定结果导出成 ``data/processed/`` 下的结果表；
* 它**不重算排放量**（主账 ``Trip_Detail!X/Y`` 的公式结果直接取缓存值），
  **不重算汇总**（``Monthly/Annual/Formal_*_Summary`` 直接取缓存值）；
* 唯一允许"算"的地方是**门禁判定链的拆解**：主账只存了合成判定
  ``Formal_Summary_Eligibility``，前端要展示"为什么 0 行能入正式汇总"，
  就得把 ``AP2`` 公式的每个谓词单独求值 —— 这属于**复现已有公式的判据**，
  不是另立一套口径，所以在本模块做、结果落表，前端只读。

与主链路的关系
--------------
* **不改** ``src/pipeline.py`` 的 E₀ 主口径与 28 条闸门；
* **不写** 主链路结果表（不碰 wide_data / baseline_carbon / final_report / kpi）；
* 碳账是 **WTW 全生命周期**口径，工程侧 E₀ 是 **TTW 尾管燃烧**口径，
  两套口径**不合并、不相减、不做趋势对比** —— 结果表里也刻意不同时放两边的总量。

用法::

    python -m src.carbon_ledger_export              # 导出到 data/processed/
    python -m src.carbon_ledger_export --no-write   # 只打印摘要，不落盘

产出（7 张，``src/schema.py`` 登记同名）::

    carbon_ledger_trips.parquet          主账 Trip_Detail 全 42 列 + 门禁谓词拆解
    carbon_ledger_evidence.parquet       证据索引 + 凭证文件存在性（实时判定）
    carbon_ledger_summary.json           账本状态条 / 月年汇总 / 正式汇总 / 口径声明
    carbon_ledger_gates.json             9 项正式汇总资格判定链
    carbon_ledger_quota.json             配额台账
    carbon_ledger_activity_readiness.json 51 字段数据就绪度
    carbon_ledger_rules.json             15 条控制规则 / 15 项用例 / 19 项压力场景 / 规则卡
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import traceback
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src import schema  # noqa: E402

LEDGER_DIR = ROOT / "carbon_ledger"
PROCESSED = ROOT / "data" / "processed"
FIELD_MAP_PATH = ROOT / "config" / "carbon_ledger_field_map.json"
CALC_FACTORS_PATH = ROOT / "config" / "calc_factors.json"

MAIN_LEDGER = LEDGER_DIR / "a4_ledger" / "carbon_ledger_main.xlsx"
QUOTA_LEDGER = LEDGER_DIR / "a4_ledger" / "quota_ledger.xlsx"
EVIDENCE_INDEX = LEDGER_DIR / "a5_evidence" / "evidence_index.xlsx"
FIELD_DICTIONARY = LEDGER_DIR / "a2_activity_data" / "activity_data_field_dictionary.xlsx"
ACTIVITY_TEMPLATE = LEDGER_DIR / "a2_activity_data" / "activity_data_template.xlsx"
RULE_CARD_MD = LEDGER_DIR / "a1_rules" / "carbon_accounting_rule_card_v2.md"

TRIPS_PATH = PROCESSED / schema.RESULT_FILES["carbon_ledger_trips"]
EVIDENCE_PATH = PROCESSED / schema.RESULT_FILES["carbon_ledger_evidence"]
SUMMARY_PATH = PROCESSED / schema.RESULT_FILES["carbon_ledger_summary"]
GATES_PATH = PROCESSED / schema.RESULT_FILES["carbon_ledger_gates"]
QUOTA_PATH = PROCESSED / schema.RESULT_FILES["carbon_ledger_quota"]
READINESS_PATH = PROCESSED / schema.RESULT_FILES["carbon_ledger_activity_readiness"]
RULES_PATH = PROCESSED / schema.RESULT_FILES["carbon_ledger_rules"]

# ---------------------------------------------------------------------------
# 正式汇总资格判定链
# ---------------------------------------------------------------------------
# 来源：carbon_ledger_main.xlsx!Trip_Detail 的 AP 列公式
#   =IF(A2="","",IF(AND(Z2="actual",AA2<>"yes",W2="factor_verified",
#       AJ2="PASS_FORMAL_PRECHECK",AN2="PASS_UNIQUE_PHYSICAL_TRIP",
#       AO2="approved",AB2<>"",AD2<>"",AE2<>""),
#       "FORMAL_APPROVED","EXCLUDED_NOT_FORMALLY_APPROVED"))
# 该 AND 里有 **9 个谓词**（docs/53 §4 记成 8 个，漏了 Source_Ledger_Row_No 非空）。
# 这里按公式原样复现 9 项，不多不少。
GATE_CONDITIONS: List[Dict[str, Any]] = [
    {
        "key": "data_status_actual",
        "column": "Data_Status",
        "label": "Data_Status = actual",
        "op": "eq",
        "expect": "actual",
        "why": "只有实测行才可能正式入账",
    },
    {
        "key": "simulated_not_yes",
        "column": "Simulated",
        "label": "Simulated ≠ yes",
        "op": "ne",
        "expect": "yes",
        "why": "模拟行不得进入正式汇总",
    },
    {
        "key": "factor_verified",
        "column": "Factor_Status",
        "label": "Factor_Status = factor_verified",
        "op": "eq",
        "expect": "factor_verified",
        "why": "排放因子须完成核验",
    },
    {
        "key": "formal_precheck",
        "column": "Validation_Status",
        "label": "Validation_Status = PASS_FORMAL_PRECHECK",
        "op": "eq",
        "expect": "PASS_FORMAL_PRECHECK",
        "why": "结构预检须通过",
    },
    {
        "key": "unique_physical_trip",
        "column": "Physical_Trip_Control",
        "label": "Physical_Trip_Control = PASS_UNIQUE_PHYSICAL_TRIP",
        "op": "eq",
        "expect": "PASS_UNIQUE_PHYSICAL_TRIP",
        "why": "物理趟次去重指纹须唯一",
    },
    {
        "key": "owner_approval",
        "column": "Owner_Approval_Status (manual)",
        "label": "Owner_Approval_Status = approved",
        "op": "eq",
        "expect": "approved",
        "why": "公式不能替负责人批准",
    },
    {
        "key": "evidence_id",
        "column": "Evidence_ID",
        "label": "Evidence_ID 非空",
        "op": "ne_blank",
        "expect": "",
        "why": "每行须挂证据编号",
    },
    {
        "key": "run_id",
        "column": "Run_ID",
        "label": "Run_ID 非空",
        "op": "ne_blank",
        "expect": "",
        "why": "每行须挂 run 版本",
    },
    {
        "key": "source_ledger_row_no",
        "column": "Source_Ledger_Row_No",
        "label": "Source_Ledger_Row_No 非空",
        "op": "ne_blank",
        "expect": "",
        "why": "来源行号须可追溯",
    },
]

#: 字段字典表头与数据**错位一列**（表头多声明了「单位」列，数据多出末列「确认状态」）。
#: 这里按**语义真序**读数据列，顺序以 activity_data_template.xlsx 的 Activity_Data
#: 表头为准（那一张是对的）。这是一个已知结构问题，记录在 docs/54 §6。
FIELD_DICT_SEMANTIC_ORDER = [
    "字段名",
    "中文含义",
    "数据类型",
    "是否必填",
    "数据来源",
    "关联财务账字段",
    "缺失处理规则",
    "校验规则",
    "备注",
    "确认状态",
]

#: 边界声明 —— 两套口径不可混算（前端常驻显示，文本也落进结果表）
BOUNDARY_STATEMENT = (
    "碳会计轨为 WTW 全生命周期口径（柴油 = 尾管 + 上游）；工程侧 E₀ 为 TTW 尾管燃烧口径。"
    "两套口径不合并、不相减、不做趋势对比，也不并排比大小 —— "
    "两边的核算范围本身也不同（碳账当前仅 12 条合成趟次，工程侧为全网趟次）。"
)


# ---------------------------------------------------------------------------
# 基础工具
# ---------------------------------------------------------------------------

def _now() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def _sha256(path: Path) -> str:
    try:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1 << 20), b""):
                digest.update(chunk)
        return digest.hexdigest()[:16]
    except OSError:
        return ""


def _text(value: Any) -> str:
    """None / NaN 一律折成空串（不编数字、不填 0）。"""
    if value is None:
        return ""
    if isinstance(value, float) and value != value:
        return ""
    return str(value).strip()


def _number(value: Any) -> Optional[float]:
    try:
        if value is None:
            return None
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number != number:
        return None
    return number


def _has_parquet_engine() -> bool:
    try:
        import pyarrow  # noqa: F401
        return True
    except Exception:  # noqa: BLE001
        try:
            import fastparquet  # noqa: F401
            return True
        except Exception:  # noqa: BLE001
            return False


def _write_json_sidecar(frame: pd.DataFrame, path: Path) -> Path:
    """写同名 JSON 侧车（人可读 + 无 Parquet 引擎时的只读回退）。"""
    sidecar = path.with_suffix(".json")
    rows = frame.astype(object).where(pd.notna(frame), None).to_dict(orient="records")
    _save_json(rows, sidecar)
    return sidecar


def _save_parquet(frame: pd.DataFrame, path: Path, tries: int = 4) -> None:
    """落 Parquet **并恒定附带同名 JSON 侧车**；缺少 Parquet 引擎时只落侧车。

    为什么要"恒定附带"：
        原实现是"有引擎写 Parquet 并**删掉**侧车，无引擎只写侧车"。
        这让**同一份代码在两种机器上产出不同的文件集** —— 谁最后跑一次，
        仓库里就留谁的形态，另一形态被删。结果表因此"不可由仓库自身复现"。

    现在的约定：
        · 有引擎 → 写 ``*.parquet``（契约格式）+ **同时**写 ``*.json`` 侧车；
        · 无引擎 → 只写 ``*.json`` 侧车，并**清掉**可能存在的过期 Parquet
          （否则读取端会优先读到上一次的旧 Parquet）。
        两种环境下 **JSON 侧车恒定存在**，读取端因此永远有一条稳定通路；
        Parquet 是否存在由 manifest 的 ``output_formats`` 如实记录。

    不伪装成 Parquet，也不改变行、列、排放数值和 simulated 状态。
    """
    import time

    path.parent.mkdir(parents=True, exist_ok=True)
    has_engine = _has_parquet_engine()
    if has_engine:
        last: Optional[Exception] = None
        for attempt in range(tries):
            try:
                frame.to_parquet(path, index=False)
                last = None
                break
            except Exception as exc:  # noqa: BLE001
                last = exc
                time.sleep(0.4 * (attempt + 1))
        # Parquet engine exists but actual write failed: fail loudly rather than silently fallback.
        if last is not None:
            raise RuntimeError("写 parquet 失败 %s：%s" % (path.name, last))
    else:
        # 没有引擎：清掉可能残留的旧 Parquet，避免读取端读到过期数据
        if path.exists():
            try:
                path.unlink()
            except OSError:
                pass
    _write_json_sidecar(frame, path)


def _save_json(payload: Any, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8")


def read_records(path: Path, sheet: str) -> List[Dict[str, Any]]:
    """读一张工作表 → 记录列表。

    * ``data_only=True``：取工作簿里**已算好的缓存值**（公式结果），不重算；
    * 兼容"首行是标题、次行才是表头"的工作表（Feasibility_Summary 就是这种）。
    """
    import openpyxl

    if not path.exists():
        raise FileNotFoundError("交付文件不存在：%s" % path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    if sheet not in workbook.sheetnames:
        raise KeyError("工作表不存在：%s!%s" % (path.name, sheet))
    worksheet = workbook[sheet]
    rows = [list(item) for item in worksheet.iter_rows(values_only=True)]
    if not rows:
        return []
    # 表头定位：跳过标题行/空行，取**第一个有 ≥2 个非空单元格的行**当表头。
    # （Feasibility_Summary 的结构是 r1 标题 + r2 空行 + r3 表头，不能按固定行号读。）
    header_index = 0
    for index, row in enumerate(rows):
        filled = sum(1 for cell in row if cell is not None and str(cell).strip())
        if filled >= 2:
            header_index = index
            break
    header = ["" if cell is None else str(cell).strip()
              for cell in rows[header_index]]
    rows = rows[header_index + 1:]
    records: List[Dict[str, Any]] = []
    for row in rows:
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        records.append({header[i]: row[i] if i < len(row) else None
                        for i in range(len(header))})
    return records


def read_sheet_header(path: Path, sheet: str) -> List[str]:
    import openpyxl

    if not path.exists():
        raise FileNotFoundError("交付文件不存在：%s" % path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook[sheet]
    return ["" if cell is None else str(cell).strip() for cell in next(
        worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]


# ---------------------------------------------------------------------------
# 主账：Trip_Detail
# ---------------------------------------------------------------------------

def build_trips() -> pd.DataFrame:
    """主账 Trip_Detail 全列 + 门禁谓词拆解。

    排放量列（``Emissions_kgCO2e`` / ``Emissions_tCO2e``）直接取工作簿缓存值，
    **不重算**。
    """
    records = read_records(MAIN_LEDGER, "Trip_Detail")
    frame = pd.DataFrame(records)
    if frame.empty:
        return frame
    for condition in GATE_CONDITIONS:
        column = condition["column"]
        if column not in frame.columns:
            frame["gate_" + condition["key"]] = False
            continue
        frame["gate_" + condition["key"]] = frame.apply(
            lambda row, cond=condition: _evaluate_condition(cond, row),
            axis=1)
    frame["gate_first_block"] = frame.apply(_first_block, axis=1)
    frame["row_formal_approved"] = (
        frame.get("Formal_Summary_Eligibility", pd.Series(dtype=object))
        .astype(str).str.strip() == "FORMAL_APPROVED")
    return frame


def _evaluate_condition(condition: Dict[str, Any], row: Any) -> bool:
    """复现 ``AP2`` 公式的一个谓词。"""
    value = _text(row.get(condition["column"]))
    op = condition["op"]
    if op == "eq":
        return value == condition["expect"]
    if op == "ne":
        return value != condition["expect"]
    if op == "ne_blank":
        return value != ""
    return False


def _first_block(row: Any) -> str:
    """返回该行**第一道**未通过的判定（没有则空）。"""
    for condition in GATE_CONDITIONS:
        if not _evaluate_condition(condition, row):
            return condition["key"]
    return ""


def build_gates(trips: pd.DataFrame) -> Dict[str, Any]:
    """9 项正式汇总资格判定链（含每项通过/阻断计数与阻断主因）。"""
    total = int(len(trips))
    rows = trips.to_dict(orient="records") if total else []
    conditions: List[Dict[str, Any]] = []
    for condition in GATE_CONDITIONS:
        flags = [_evaluate_condition(condition, row) for row in rows]
        passed = int(sum(1 for flag in flags if flag))
        blocked = total - passed
        observed = Counter(
            _text(row.get(condition["column"])) or "(空)"
            for row, ok in zip(rows, flags) if not ok)
        conditions.append({
            "key": condition["key"],
            "label": condition["label"],
            "column": condition["column"],
            "why": condition["why"],
            "passed": passed,
            "blocked": blocked,
            "status": "pass" if blocked == 0 else ("fail" if passed == 0 else "partial"),
            "block_reasons": [
                {"value": value, "count": int(count)}
                for value, count in observed.most_common(4)],
        })
    first_block = Counter(
        _first_block(row) for row in rows) if rows else Counter()
    eligibility = Counter(
        _text(row.get("Formal_Summary_Eligibility")) or "(空)" for row in rows)
    approved = int(sum(
        1 for row in rows
        if _text(row.get("Formal_Summary_Eligibility")) == "FORMAL_APPROVED"))
    return {
        "formula_source": "carbon_ledger_main.xlsx!Trip_Detail!AP2",
        "predicate_count": len(GATE_CONDITIONS),
        "total_rows": total,
        "approved_rows": approved,
        "conditions": conditions,
        "first_block": [
            {"key": key, "count": int(count)}
            for key, count in first_block.most_common()
        ],
        "eligibility_distribution": [
            {"value": value, "count": int(count)}
            for value, count in eligibility.most_common()
        ],
        "boundary": BOUNDARY_STATEMENT,
    }


# ---------------------------------------------------------------------------
# 证据链
# ---------------------------------------------------------------------------

def _resolve_evidence_path(raw: str) -> Tuple[str, Optional[str]]:
    """在仓库里找凭证文件；返回 (相对路径, 命中的相对根目录)。"""
    if not raw:
        return "", None
    candidates = [
        (LEDGER_DIR / "a5_evidence" / raw, "carbon_ledger/a5_evidence"),
        (LEDGER_DIR / raw, "carbon_ledger"),
        (ROOT / raw, "仓库根"),
    ]
    for path, base in candidates:
        if path.exists():
            return str(path.relative_to(ROOT)).replace("\\", "/"), base
    return raw, None


def build_evidence() -> pd.DataFrame:
    """证据索引 + **实时**判定凭证文件是否真的存在。

    原表 ``Original_Document_Status`` 是交付方的自述；本页要回答"凭证在哪"，
    所以适配器额外算一个 ``evidence_file_exists``（文件系统事实）。
    """
    records = read_records(EVIDENCE_INDEX, "Evidence_Index")
    frame = pd.DataFrame(records)
    if frame.empty:
        return frame
    exists: List[bool] = []
    resolved: List[str] = []
    row_level: List[bool] = []
    status_bucket: List[str] = []
    placeholder_flags: List[bool] = []
    source_document_flags: List[bool] = []
    original_activity_flags: List[bool] = []
    artifact_kinds: List[str] = []
    for row in frame.to_dict(orient="records"):
        raw = _text(row.get("Evidence_File_Path"))
        relative, _base = _resolve_evidence_path(raw)
        resolved.append(relative)
        artifact_exists = bool(_base)
        exists.append(artifact_exists)
        ledger_row = _text(row.get("Ledger_Row_ID"))
        is_row = bool(ledger_row) and ledger_row.upper() != "N/A"
        row_level.append(is_row)
        status = _text(row.get("Original_Document_Status"))
        if not status:
            status_bucket.append("blank")
        elif status.upper().startswith("NOT PROVIDED"):
            status_bucket.append("not_provided")
        else:
            status_bucket.append("provided")

        artifact_path = (ROOT / relative) if relative else None
        artifact_text = ""
        if artifact_exists and artifact_path is not None:
            try:
                artifact_text = artifact_path.read_text(encoding="utf-8", errors="replace")[:4000]
            except (OSError, UnicodeError):
                artifact_text = ""
        is_placeholder = bool(
            "SIMULATED PLACEHOLDER ONLY" in artifact_text
            or raw.replace("\\", "/").startswith("simulated_evidence/")
            or (_text(row.get("Simulated")).lower() == "yes" and artifact_exists)
        )
        source_document_available = artifact_exists and not is_placeholder
        original_activity_available = (
            source_document_available and is_row
            and _text(row.get("Simulated")).lower() != "yes"
            and not status.upper().startswith("NOT PROVIDED")
        )
        placeholder_flags.append(is_placeholder)
        source_document_flags.append(source_document_available)
        original_activity_flags.append(original_activity_available)
        if is_placeholder:
            artifact_kinds.append("simulated_placeholder_not_original_evidence")
        elif original_activity_available:
            artifact_kinds.append("original_activity_evidence_file")
        elif source_document_available:
            artifact_kinds.append("non_placeholder_source_document")
        else:
            artifact_kinds.append("missing_file")
    frame["evidence_file_exists"] = exists
    frame["evidence_file_resolved"] = resolved
    frame["is_row_level"] = row_level
    frame["status_bucket"] = status_bucket
    frame["evidence_is_simulated_placeholder"] = placeholder_flags
    frame["source_reference_file_available"] = source_document_flags
    frame["original_activity_evidence_available"] = original_activity_flags
    frame["evidence_artifact_kind"] = artifact_kinds
    return frame


# ---------------------------------------------------------------------------
# 配额台账
# ---------------------------------------------------------------------------

def build_quota() -> Dict[str, Any]:
    """配额台账（单条记录）+ 台账说明。**字段原文照搬，不改写**。"""
    records = read_records(QUOTA_LEDGER, "Quota_Ledger")
    if not records:
        return {}
    row = records[0]
    instructions: List[Dict[str, Any]] = []
    try:
        for item in read_records(QUOTA_LEDGER, "Quota_Instructions"):
            instructions.append({"项目": _text(item.get("项目")),
                                 "说明": _text(item.get("说明"))})
    except KeyError:
        instructions = []
    return {
        "quota_record_id": _text(row.get("Quota_Record_ID")),
        "period": _text(row.get("Period")),
        "opening_quota_t": _number(row.get("Opening_Quota_tCO2e")),
        "compliance_used_t": _number(row.get("Compliance_Used_tCO2e")),
        "purchased_quota_t": _number(row.get("Purchased_Quota_tCO2e")),
        "available_quota_t": _number(row.get("Available_Quota_tCO2e")),
        "closing_balance_t": _number(row.get("Closing_Balance_tCO2e")),
        "shortfall_t": _number(row.get("Shortfall_tCO2e")),
        "data_status": _text(row.get("Data_Status")),
        "approval_status": _text(row.get("Approval_Status")),
        "source_document": _text(row.get("Source_Document")),
        "notes": _text(row.get("Notes")),
        "instructions": instructions,
        "raw_fields": {key: ("" if value is None else value)
                       for key, value in row.items()},
    }


# ---------------------------------------------------------------------------
# 活动数据 51 字段就绪度
# ---------------------------------------------------------------------------

def build_readiness() -> Dict[str, Any]:
    """51 字段的数据就绪度：今天能不能填、谁来填。

    分类来自 ``config/carbon_ledger_field_map.json``（本仓库维护的映射），
    字段清单来自交付方的字段字典。
    """
    field_map: Dict[str, Any] = {}
    categories: List[Dict[str, Any]] = []
    if FIELD_MAP_PATH.exists():
        payload = json.loads(FIELD_MAP_PATH.read_text(encoding="utf-8"))
        field_map = payload.get("fields") or {}
        categories = payload.get("categories") or []
    category_labels = {item.get("id"): item.get("label")
                       for item in categories}

    records = read_records(FIELD_DICTIONARY, "Field_Dictionary")
    template_fields: List[str] = []
    template_filled = 0
    if ACTIVITY_TEMPLATE.exists():
        template_fields = read_sheet_header(ACTIVITY_TEMPLATE, "Activity_Data")
        try:
            template_rows = read_records(ACTIVITY_TEMPLATE, "Activity_Data")
            template_filled = len(template_rows)
        except Exception:  # noqa: BLE001
            template_filled = 0

    fields: List[Dict[str, Any]] = []
    seen: List[str] = []
    for record in records:
        # 表头与数据错位一列 → 按语义真序重排
        item = dict(zip(FIELD_DICT_SEMANTIC_ORDER, list(record.values())))
        name = _text(item.get("字段名"))
        if not name:
            continue
        seen.append(name)
        mapping = field_map.get(name) or {}
        category = str(mapping.get("category") or "external")
        fields.append({
            "field": name,
            "label_zh": _text(item.get("中文含义")),
            "data_type": _text(item.get("数据类型")),
            "required": _text(item.get("是否必填")),
            "source_text": _text(item.get("数据来源")),
            "finance_field": _text(item.get("关联财务账字段")),
            "missing_rule": _text(item.get("缺失处理规则")),
            "validation_rule": _text(item.get("校验规则")),
            "remark": _text(item.get("备注")),
            "confirm_status": _text(item.get("确认状态")),
            "category": category,
            "category_label": category_labels.get(category, category),
            "note": _text(mapping.get("note")),
        })

    missing_in_map = [name for name in seen if name not in field_map]
    extra_in_map = [name for name in field_map if name not in seen]
    counts = Counter(item["category"] for item in fields)
    return {
        "field_count": len(fields),
        "categories": categories,
        "counts": [
            {"category": item.get("id"),
             "label": category_labels.get(item.get("id"), item.get("id")),
             "desc": item.get("desc"),
             "count": int(counts.get(item.get("id"), 0))}
            for item in categories
        ],
        "fields": fields,
        "template_field_count": len(template_fields),
        "template_filled_rows": template_filled,
        "unmapped_fields": missing_in_map,
        "unused_map_entries": extra_in_map,
        "note": ("就绪度是**本仓库的映射判断**，不是交付方的结论；"
                 "「待财务线补」「待外部资料」的字段当前确实没有数据。"),
    }


# ---------------------------------------------------------------------------
# 规则与自证
# ---------------------------------------------------------------------------

def build_rules() -> Dict[str, Any]:
    """控制规则 / 基础用例 / 压力场景 / 可行性自查 / 规则卡原文。"""
    def _table(sheet: str) -> List[Dict[str, Any]]:
        try:
            return read_records(MAIN_LEDGER, sheet)
        except Exception:  # noqa: BLE001
            return []

    validation = _table("Validation_Rules")
    tests = _table("Test_Cases")
    pressure = _table("Pressure_Test_Cases")
    feasibility = _table("Feasibility_Summary")
    pending = _table("Pending_4_Vehicle_Data")
    card = ""
    if RULE_CARD_MD.exists():
        card = RULE_CARD_MD.read_text(encoding="utf-8")
        if len(card) > 120_000:
            card = card[:120_000]

    def _passed(rows: List[Dict[str, Any]]) -> Tuple[int, int]:
        passed = 0
        total = 0
        for row in rows:
            joined = " ".join(_text(value) for value in row.values()).upper()
            if "PASS" not in joined and "FAIL" not in joined:
                continue
            total += 1
            if "PASS" in joined and "FAIL" not in joined:
                passed += 1
        return passed, total

    pressure_passed, pressure_total = _passed(pressure)
    self_check_path = LEDGER_DIR / "a6_verification" / "self_check.py"
    self_check_result_path = LEDGER_DIR / "a6_verification" / "self_check_results.json"
    self_check_result: Dict[str, Any] = {}
    if self_check_result_path.exists():
        try:
            self_check_result = json.loads(self_check_result_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            self_check_result = {}
    placeholder_files = sorted((LEDGER_DIR / "a5_evidence" / "simulated_evidence").glob("*.txt"))
    return {
        "validation_rules": validation,
        "validation_rule_count": len(validation),
        "test_cases": tests,
        "test_case_count": len(tests),
        "pressure_test_cases": pressure,
        "pressure_case_count": len(pressure),
        "pressure_passed": pressure_passed,
        "pressure_checked": pressure_total,
        "feasibility_summary": feasibility,
        "pending_items": pending,
        "rule_card_markdown": card,
        "rule_card_file": str(RULE_CARD_MD.relative_to(ROOT)).replace("\\", "/"),
        "self_check": {
            "script_available": self_check_path.exists(),
            "result_available": bool(self_check_result),
            "passed": int(self_check_result.get("passed") or 0),
            "failed": int(self_check_result.get("failed") or 0),
            "result_file": str(self_check_result_path.relative_to(ROOT)).replace("\\", "/"),
        },
        "simulated_placeholder_file_count": len(placeholder_files),
        "note": ("适配器导出的是工作簿缓存结果，不重新计算排放。"
                 "本仓库随交付提供的 simulated_evidence 是链接测试占位文件，不是真实发票或设备记录；"
                 "必须区分原始活动凭证与占位文件，不能把文件存在等同于原始凭证已提供。"),
    }


# ---------------------------------------------------------------------------
# 汇总
# ---------------------------------------------------------------------------

def _load_calc_factors() -> Dict[str, Any]:
    try:
        return json.loads(CALC_FACTORS_PATH.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {}


def build_summary(trips: pd.DataFrame, gates: Dict[str, Any]) -> Dict[str, Any]:
    """账本状态条所需的一切（来源、口径、月年汇总、正式汇总）。"""
    def _table(sheet: str) -> List[Dict[str, Any]]:
        try:
            return read_records(MAIN_LEDGER, sheet)
        except Exception:  # noqa: BLE001
            return []

    monthly = _table("Monthly_Summary")
    annual = _table("Annual_Summary")
    formal_monthly = _table("Formal_Monthly_Summary")
    formal_annual = _table("Formal_Annual_Summary")

    total_rows = int(len(trips))
    emissions_t = 0.0
    simulated_rows = 0
    actual_rows = 0
    included_rows = 0
    included_t = 0.0
    ef_values: Counter = Counter()
    factor_status: Counter = Counter()
    for row in trips.to_dict(orient="records"):
        value = _number(row.get("Emissions_tCO2e"))
        if value is not None:
            emissions_t += value
        status = _text(row.get("Data_Status")).lower()
        if status == "simulated":
            simulated_rows += 1
        elif status == "actual":
            actual_rows += 1
        if _text(row.get("Include_In_Summary")).upper() == "Y":
            included_rows += 1
            if value is not None:
                included_t += value
        ef = _number(row.get("EF_Value"))
        unit = _text(row.get("EF_Unit"))
        if ef is not None:
            ef_values["%s kgCO2e/%s" % (_trim(ef), unit)] += 1
        fs = _text(row.get("Factor_Status"))
        if fs:
            factor_status[fs] += 1

    formal_row_count = 0
    formal_total = 0.0
    formal_gate_status = ""
    formal_note = ""
    for row in formal_annual:
        formal_row_count += int(_number(row.get("Approved_Row_Count")) or 0)
        formal_total += float(_number(row.get("Approved_Emissions_tCO2e")) or 0.0)
        if not formal_gate_status:
            formal_gate_status = _text(row.get("Gate_Status"))
        if not formal_note:
            formal_note = _text(row.get("Note"))

    factors = _load_calc_factors()
    co2 = factors.get("co2_factors") or {}
    ttw = float((co2.get("diesel_ttw_kg_per_l") or {}).get("value") or 0.0)
    wtt = float((co2.get("diesel_wtt_kg_per_l") or {}).get("value") or 0.0)
    wtw_total = round(ttw + wtt, 6)
    ledger_ef = _ledger_diesel_ef(trips)
    ef_agrees = bool(ledger_ef) and abs(ledger_ef - wtw_total) < 1e-6

    return {
        "generated_at": _now(),
        "ledger_file": str(MAIN_LEDGER.relative_to(ROOT)).replace("\\", "/"),
        "ledger_sha256_16": _sha256(MAIN_LEDGER),
        "workbook_sheet_count": 13,
        "trip_rows": total_rows,
        "total_emissions_t": round(emissions_t, 9),
        "simulated_rows": simulated_rows,
        "actual_rows": actual_rows,
        "include_in_summary_rows": included_rows,
        "include_in_summary_t": round(included_t, 9),
        "is_simulated": simulated_rows > 0,
        "ledger_status": ("未获正式批准" if gates.get("approved_rows", 0) == 0
                          else "存在正式批准行"),
        "monthly": [
            {"month": _text(row.get("Month")),
             "rows": int(_number(row.get("Included_Row_Count")) or 0),
             "total_t": _number(row.get("Total_Emissions_tCO2e")),
             "data_state": _text(row.get("Data_State")),
             "report_status": _text(row.get("Report_Status"))}
            for row in monthly
        ],
        "annual": [
            {"year": _text(row.get("Year")),
             "rows": int(_number(row.get("Included_Row_Count")) or 0),
             "total_t": _number(row.get("Total_Emissions_tCO2e")),
             "data_state": _text(row.get("Data_State")),
             "report_status": _text(row.get("Report_Status"))}
            for row in annual
        ],
        "formal": {
            "annual_rows": len(formal_annual),
            "monthly_rows": len(formal_monthly),
            "approved_row_count": formal_row_count,
            "approved_total_t": formal_total,
            "gate_status": formal_gate_status,
            "note": formal_note,
            "meaning": ("0 行 = 没有合格正式行，**不等于实际排放为零**"
                        if formal_row_count == 0 else "存在合格正式行"),
        },
        "factor": {
            "ledger_ef_values": [{"value": key, "rows": int(count)}
                                 for key, count in ef_values.most_common()],
            "factor_status": [{"value": key, "rows": int(count)}
                              for key, count in factor_status.most_common()],
            "basis": "WTW",
            "basis_label": "WTW 全生命周期（尾管燃烧 + 上游）",
            "engineering_basis": "TTW",
            "engineering_basis_label": "TTW 尾管燃烧（工程侧 E₀ 口径）",
            "mix_allowed": False,
            "decomposition": {
                "diesel_ttw_kg_per_l": ttw,
                "diesel_wtt_kg_per_l": wtt,
                "diesel_wtw_kg_per_l": wtw_total,
                "grid_kg_per_kwh": float(
                    (co2.get("grid_wtw_kg_per_kwh") or {}).get("value") or 0.0),
                "source": "config/calc_factors.json",
            },
            "ledger_ef_matches_calc_factors": ef_agrees,
        },
        "boundary": BOUNDARY_STATEMENT,
        "source_files": [
            str(path.relative_to(ROOT)).replace("\\", "/")
            for path in (MAIN_LEDGER, QUOTA_LEDGER, EVIDENCE_INDEX,
                         FIELD_DICTIONARY, ACTIVITY_TEMPLATE)
            if path.exists()
        ],
    }


def _trim(number: float) -> str:
    text = ("%.4f" % number).rstrip("0").rstrip(".")
    return text or "0"


def _ledger_diesel_ef(trips: pd.DataFrame) -> Optional[float]:
    for row in trips.to_dict(orient="records"):
        if _text(row.get("Activity_Unit")) == "L" and _number(row.get("EF_Value")):
            return _number(row.get("EF_Value"))
    return None


# ---------------------------------------------------------------------------
# 导出
# ---------------------------------------------------------------------------

def export(write: bool = True, out_dir: Optional[Path] = None) -> Dict[str, Any]:
    """跑一遍适配器，返回摘要 dict。

    ``out_dir`` 给了就写到别处（测试用临时目录），**避免测试把
    ``data/processed/`` 里已入库的结果表改成脏工作树**。
    """
    trips = build_trips()
    if trips.empty:
        raise RuntimeError("主账 Trip_Detail 读不到数据行：%s" % MAIN_LEDGER)
    gates = build_gates(trips)
    evidence = build_evidence()
    quota = build_quota()
    readiness = build_readiness()
    rules = build_rules()
    summary = build_summary(trips, gates)

    base = Path(out_dir) if out_dir else PROCESSED
    outputs = {
        base / schema.RESULT_FILES["carbon_ledger_trips"]: trips,
        base / schema.RESULT_FILES["carbon_ledger_evidence"]: evidence,
        base / schema.RESULT_FILES["carbon_ledger_summary"]: summary,
        base / schema.RESULT_FILES["carbon_ledger_gates"]: gates,
        base / schema.RESULT_FILES["carbon_ledger_quota"]: quota,
        base / schema.RESULT_FILES["carbon_ledger_activity_readiness"]: readiness,
        base / schema.RESULT_FILES["carbon_ledger_rules"]: rules,
    }
    has_engine = _has_parquet_engine()
    output_formats = {
        path.name: ("parquet_with_json_sidecar" if isinstance(payload, pd.DataFrame)
                    and has_engine else ("json_sidecar_no_parquet_engine"
                    if isinstance(payload, pd.DataFrame) else "json"))
        for path, payload in outputs.items()}
    generated_at = _now()
    if write:
        for path, payload in outputs.items():
            if isinstance(payload, pd.DataFrame):
                _save_parquet(payload, path)
            else:
                _save_json(payload, path)
        _save_json({
            "generated_at": generated_at,
            "output_formats": output_formats,
            "format_notice": "JSON sidecars are explicit local compatibility outputs, not Parquet files; standard deployment with pyarrow/fastparquet emits Parquet.",
            "simulated_rows": int(summary.get("simulated_rows") or 0),
            "actual_rows": int(summary.get("actual_rows") or 0),
            "approved_rows": int(gates.get("approved_rows") or 0),
            "formal_gate_status": summary.get("formal", {}).get("gate_status", ""),
        }, base / "carbon_ledger_export_manifest.json")

    return {
        "generated_at": generated_at,
        "written": write,
        "out_dir": str(base),
        "outputs": {path.name: (len(payload) if isinstance(payload, pd.DataFrame)
                                else 1) for path, payload in outputs.items()},
        "output_formats": output_formats,
        "trip_rows": int(len(trips)),
        "total_emissions_t": summary["total_emissions_t"],
        "approved_rows": gates["approved_rows"],
        "formal_gate_status": summary["formal"]["gate_status"],
        "evidence_rows": int(len(evidence)),
        "evidence_files_present": int(evidence["evidence_file_exists"].sum())
        if not evidence.empty else 0,
        "quota_shortfall_t": quota.get("shortfall_t"),
        "activity_fields": readiness["field_count"],
        "validation_rules": rules["validation_rule_count"],
        "pressure_cases": rules["pressure_case_count"],
    }


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="碳账交付 → 前端只读结果表（适配器）")
    parser.add_argument("--no-write", action="store_true",
                        help="只打印摘要，不写 data/processed/")
    args = parser.parse_args(argv)
    try:
        report = export(write=not args.no_write)
    except Exception as exc:  # noqa: BLE001
        print("[carbon_ledger_export] 失败：%s" % exc)
        traceback.print_exc()
        return 1
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
